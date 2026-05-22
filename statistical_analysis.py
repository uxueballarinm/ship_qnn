# %% [markdown]
# # Ultimate Interactive QNN/QELM Statistical Engine
# Upgraded parameter schema from 'agg_override' to an adaptive 'strategy' configuration.

# %%
import os
import warnings
import numpy as np
import pandas as pd
import seaborn as sns
import matplotlib.pyplot as plt
import statsmodels.api as sm
from scipy import stats
from statsmodels.graphics.gofplots import qqplot
from statsmodels.stats.multitest import multipletests
from statsmodels.formula.api import ols
import openpyxl
from openpyxl.styles import PatternFill
import openpyxl.utils as utils
from openpyxl.formatting.rule import CellIsRule

# STYLE CONFIGURATION
warnings.filterwarnings("ignore")
plt.rcParams.update({
    "font.family": "serif",
    "font.serif": ["Cambria", "Times New Roman"],
    "axes.titleweight": "bold",
    "font.size": 11
})

class InteractiveQNNAnalyzer:
    def __init__(self, file_path, sheet_name=0, save_enabled=True):
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        self.save_enabled = save_enabled
        self.file_path = file_path
        self.sheet_name = sheet_name
        self.df = pd.read_excel(file_path, sheet_name=sheet_name, engine='openpyxl')
        self.my_colors = ["#4DB6AC", "#00695C", "#004D40", "#B2DFDB", "#80CBC4"]
        self.line_color = "#002B2E" 
        
        if self.save_enabled:
            base_dir = os.path.dirname(os.path.abspath(file_path))
            self.out_dir = os.path.join(base_dir, "statistical_analysis_results")
            os.makedirs(self.out_dir, exist_ok=True)
            
            self.report_path = os.path.join(self.out_dir, "analysis_report.xlsx")
            
            with pd.ExcelWriter(self.report_path, engine='openpyxl', mode='w') as writer:
                pd.DataFrame({"Status": ["Execution Active"]}).to_excel(writer, sheet_name="Summary_Overview", index=False)
        else:
            self.report_path, self.out_dir = None, None

        self._preprocess()
        self._create_config_id()
    
    def _preprocess(self):
        if 'map' in self.df.columns:
            self.df['map'] = self.df['map'].astype(str).str.replace(" ", "")
        categorical_cols = ['reps', 'encoding', 'ansatz', 'entangle', 'head_number', 'initialization']
        for col in categorical_cols:
            if col in self.df.columns:
                self.df[col] = self.df[col].astype(str).str.strip()

    def _create_config_id(self):
        group_keys = ["features", "targets", "window_size", "horizon", "model", "head_number", "encoding", "ansatz", "entangle", "reps", "map", "initialization"]
        existing_keys = [k for k in group_keys if k in self.df.columns]
        
        self.df['Full_Config_Long'] = self.df[existing_keys].astype(str).agg(' | '.join, axis=1)
        unique_configs = self.df['Full_Config_Long'].unique()
        mapping_dict = {long: f"Cfg_{i+1}" for i, long in enumerate(unique_configs)}
        self.df['configuration'] = self.df['Full_Config_Long'].map(mapping_dict)
        
        if self.save_enabled:
            mapping_df = self.df[['configuration'] + existing_keys].drop_duplicates().sort_values('configuration')
            mapping_df.set_index('configuration', inplace=True)
            self._safe_append_sheet(mapping_df, "Configuration_Mapping")
            print("🗺️ Configuration Mapping key saved to sheet: 'Configuration_Mapping'")

    def _apply_filters(self, data_df, filters_dict):
        if not filters_dict:
            return data_df
        filtered_df = data_df.copy()
        for key, val in filters_dict.items():
            if key in filtered_df.columns:
                filtered_df = filtered_df[filtered_df[key].astype(str) == str(val)]
            else:
                print(f"Warning: Filter column '{key}' not found in dataset.")
        return filtered_df

    def _safe_append_sheet(self, dataframe, sheet_name, is_matrix=False):
        if not self.save_enabled or self.report_path is None:
            return
        for char in [':', '?', '*', '/', '\\', '[', ']']:
            sheet_name = sheet_name.replace(char, '')
        sheet_name = sheet_name[:30]
        
        with pd.ExcelWriter(self.report_path, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
            dataframe.to_excel(writer, sheet_name=sheet_name)
            
            if is_matrix:
                ws = writer.sheets[sheet_name]
                green = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                red = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                cell_range = f"B2:{utils.get_column_letter(ws.max_column)}{ws.max_row}"
                ws.conditional_formatting.add(cell_range, CellIsRule(operator='lessThan', formula=['0.05'], fill=green))
                ws.conditional_formatting.add(cell_range, CellIsRule(operator='greaterThanOrEqual', formula=['0.05'], fill=red))

    def check_normality_full(self, data):
        N = len(data)
        if N < 3: return False
        if N < 50:
            _, p = stats.shapiro(data)
        else:
            _, p = stats.normaltest(data)
        return p > 0.05

    def calculate_hedges_g(self, group1, group2):
        n1, n2 = len(group1), len(group2)
        if n1 < 2 or n2 < 2: return 0
        mean1, mean2 = np.mean(group1), np.mean(group2)
        var1, var2 = np.var(group1, ddof=1), np.var(group2, ddof=1)
        pooled_std = np.sqrt(((n1 - 1) * var1 + (n2 - 1) * var2) / (n1 + n2 - 2))
        if pooled_std == 0: return 0
        d = (mean1 - mean2) / pooled_std
        return d * (1 - (3 / (4 * (n1 + n2) - 9)))

    def run_study(self, group_by='configuration', metric='Val Global Open R2', strategy=None, filters=None):
        """Compares best vs rest using either 'mean', 'median', 'best' or automatic normality resolution."""
        print(f"\n{'='*80}\nSTUDY: {metric} grouped by '{group_by}'\nFilters Applied: {filters}\n{'='*80}")
        
        self.df[metric] = pd.to_numeric(self.df[metric], errors='coerce')
        clean_df = self.df.dropna(subset=[metric]).copy()
        clean_df = self._apply_filters(clean_df, filters)
        
        if clean_df.empty:
            print("❌ Error: Filter constraints left 0 data rows.")
            return None

        is_error_metric = any(x in metric.upper() for x in ["MSE", "LOSS", "MAE"])
        unique_groups = clean_df[group_by].unique()
        
        norm_results = {idx: self.check_normality_full(clean_df[clean_df[group_by]==idx][metric]) for idx in unique_groups}
        pct_normal = sum(norm_results.values()) / len(norm_results) if norm_results else 0
        use_parametric = pct_normal > 0.95
        
        # Strategy selection matrix mapping logic
        if strategy == 'best':
            agg_metric = 'min' if is_error_metric else 'max'
        elif strategy == 'mean':
            agg_metric = 'mean'
        elif strategy == 'median':
            agg_metric = 'median'
        else: # strategy is None
            agg_metric = 'mean' if use_parametric else 'median'
        
        print(f">> Normality Pass Rate: {pct_normal:.1%} | Strategy Selected: {agg_metric.upper()}")

        stats_table = clean_df.groupby(group_by)[metric].agg(['mean', 'median', 'std', 'count', 'max', 'min']).copy()
        stats_table = stats_table.sort_values(agg_metric, ascending=is_error_metric)

        if len(stats_table) > 1:
            best_idx = stats_table.index[0]
            d_best_df = clean_df[clean_df[group_by] == best_idx].sort_values('run')
            raw_p_values, effect_sizes, evaluated_indices = [], [], []

            for other_idx in stats_table.index[1:]:
                d_other_df = clean_df[clean_df[group_by] == other_idx].sort_values('run')
                merged = pd.merge(d_best_df, d_other_df, on='run', suffixes=('_b', '_o'))
                d_best = merged[f'{metric}_b'].values
                d_other = merged[f'{metric}_o'].values
                
                if len(d_best) >= 3:
                    _, p = stats.ttest_rel(d_best, d_other) if use_parametric else stats.wilcoxon(d_best, d_other, alternative='two-sided')
                    raw_p_values.append(p)
                    effect_sizes.append(self.calculate_hedges_g(d_best, d_other))
                    evaluated_indices.append(other_idx)
            
            if raw_p_values:
                reject, p_adjusted, _, _ = multipletests(raw_p_values, alpha=0.05, method='fdr_bh')
                adj_p_map = dict(zip(evaluated_indices, p_adjusted))
                sig_map = dict(zip(evaluated_indices, ["YES" if r else "NO" for r in reject]))
                
                def interpret_g(g):
                    if pd.isna(g): return "N/A"
                    g_abs = abs(g)
                    if g_abs < 0.2: return "Negligible"
                    if g_abs < 0.5: return "Small"
                    if g_abs < 0.8: return "Medium"
                    return "Large"
                
                g_map = dict(zip(evaluated_indices, [interpret_g(val) for val in effect_sizes]))
                
                stats_table['Hedges_g_vs_Best'] = stats_table.index.map(lambda x: "N/A (BEST)" if x == best_idx else g_map.get(x, "N/A"))
                stats_table['p-value (Adjusted)'] = stats_table.index.map(lambda x: 1.0 if x == best_idx else adj_p_map.get(x, np.nan))
                stats_table['Is_Significant'] = stats_table.index.map(lambda x: "N/A (BEST)" if x == best_idx else sig_map.get(x, "N/A"))
        
        suff = "_".join([f"{k[:3]}={v}" for k, v in filters.items()]) if filters else "Raw"
        self._safe_append_sheet(stats_table, f"BvR_{group_by[:5]}_{suff}")
            
        self._visualize_study(clean_df, stats_table, group_by, metric, agg_metric)
        return stats_table

    def _visualize_study(self, clean_df, stats_table, group_by, metric, agg_metric):
        top_groups = stats_table.index[:10].tolist()
        plot_df = clean_df[clean_df[group_by].isin(top_groups)].copy()
        plot_df[group_by] = pd.Categorical(plot_df[group_by], categories=top_groups, ordered=True)
        
        fig, axes = plt.subplots(1, 2, figsize=(14, 5))
        sns.boxplot(x=group_by, y=metric, data=plot_df, palette=self.my_colors, ax=axes[0], showmeans=True if agg_metric in ['mean', 'max', 'min'] else False)
        sns.swarmplot(x=group_by, y=metric, data=plot_df, color=self.line_color, alpha=0.4, ax=axes[0])
        axes[0].tick_params(axis='x', rotation=30)
        axes[0].set_title(f"Top Performers Matrix ({metric})")
        
        qqplot(clean_df[clean_df[group_by] == stats_table.index[0]][metric], line='s', ax=axes[1])
        axes[1].set_title("Normality Diagnostic Check")
        plt.tight_layout()
        plt.show()

    def run_pairwise_analysis(self, group_by='configuration', metric='Val Global Open R2', limit=50, strategy=None, filters=None):
        print(f"\n{'='*80}\nPAIRWISE ANALYSIS: {metric} grouped by '{group_by}'")
        print(f"Limit Configs: {limit if limit is not None else 'ALL CONFIGURATIONS'}")
        print(f"Filters Applied: {filters}\n{'='*80}")
        
        self.df[metric] = pd.to_numeric(self.df[metric], errors='coerce')
        clean_df = self.df.dropna(subset=[metric]).copy()
        clean_df = self._apply_filters(clean_df, filters)
        
        if clean_df.empty:
            print("❌ Error: Filter criteria returned no data entries.")
            return None
            
        is_error_metric = any(x in metric.upper() for x in ["MSE", "LOSS", "MAE"])
        
        if strategy == 'best':
            agg_metric = 'min' if is_error_metric else 'max'
        elif strategy == 'mean':
            agg_metric = 'mean'
        elif strategy == 'median':
            agg_metric = 'median'
        else: # strategy is None
            unique_groups = clean_df[group_by].unique()
            norm_results = {idx: self.check_normality_full(clean_df[clean_df[group_by]==idx][metric]) for idx in unique_groups}
            pct_normal = sum(norm_results.values()) / len(norm_results) if norm_results else 0
            use_parametric = pct_normal > 0.95
            agg_metric = 'mean' if use_parametric else 'median'
            
        order_idx = clean_df.groupby(group_by)[metric].agg(agg_metric).sort_values(ascending=is_error_metric).index
        
        if limit is None:
            indices = order_idx.tolist()
        else:
            indices = order_idx[:limit].tolist()
            
        print(f"-> Executing analysis matrix across {len(indices)} unique configurations sorted by {agg_metric.upper()}...")
        
        matrix = pd.DataFrame(1.0, index=indices, columns=indices)
        comparison_records = []
        
        for i_idx, i in enumerate(indices):
            for j_idx, j in enumerate(indices):
                if i_idx < j_idx:
                    d1_df = clean_df[clean_df[group_by] == i].sort_values('run')
                    d2_df = clean_df[clean_df[group_by] == j].sort_values('run')
                    merged = pd.merge(d1_df, d2_df, on='run', suffixes=('_1', '_2'))
                    d1, d2 = merged[f'{metric}_1'].values, merged[f'{metric}_2'].values
                    
                    if len(d1) >= 3:
                        _, p = stats.wilcoxon(d1, d2, alternative='two-sided')
                        comparison_records.append({'i': i, 'j': j, 'p_raw': p})
        
        if comparison_records:
            raw_p_list = [comp['p_raw'] for comp in comparison_records]
            _, p_adjusted, _, _ = multipletests(raw_p_list, alpha=0.05, method='fdr_bh')
            for idx, comp in enumerate(comparison_records):
                adj_val = p_adjusted[idx]
                matrix.loc[comp['i'], comp['j']] = adj_val
                matrix.loc[comp['j'], comp['i']] = adj_val

        self._safe_append_sheet(matrix, f"PW_{group_by[:10]}", is_matrix=True)
            
        if len(indices) <= 60:
            plt.figure(figsize=(8, 6))
            sns.heatmap(matrix.astype(float), cmap="viridis_r", vmin=0.0, vmax=0.05, annot=False, cbar_kws={'label': 'Adjusted p-value'})
            plt.title(f"Head-to-Head Pairwise Heatmap ({metric})")
            plt.show()
        else:
            print("⚠️ Heatmap rendering bypassed: Configuration count too high for an image layout.")
        
        print("\n>>> Interactive Colored Output View (Green = Significant, Red = Not Significant):")
        styled_matrix = matrix.style.background_gradient(cmap='RdYlGn_r', vmin=0.0, vmax=0.05, axis=None).format("{:.4f}")
        return styled_matrix

    def run_interaction_study(self, factor_x='map', factor_hue='initialization', metric='Val Global Open R2', limit=5, strategy=None, filters=None):
        print(f"\n{'='*80}\nINTERACTION ENGINE: {factor_x} x {factor_hue} on {metric}")
        print(f"Visual Plot Filter: Isolation bound to top {limit if limit else 'ALL'} options")
        print(f"Filters Applied: {filters}\n{'='*80}")
        
        self.df[metric] = pd.to_numeric(self.df[metric], errors='coerce')
        clean_df = self.df.dropna(subset=[metric]).copy()
        clean_df = self._apply_filters(clean_df, filters)
        
        if clean_df.empty:
            print("❌ Error: Filter parameters completely stripped target rows.")
            return None

        unique_x = clean_df[factor_x].nunique()
        unique_hue = clean_df[factor_hue].nunique()
        if unique_x < 2 or unique_hue < 2:
            print(f"🛑 Cannot calculate interaction model: {factor_x} or {factor_hue} lacks categorical variation.")
            return None

        # --- STEP 1: CALCULATE STATISTICAL TESTS RIGOROUSLY ACROSS ALL VALID DATA ---
        groups = [g[metric].values for n, g in clean_df.groupby([factor_x, factor_hue])]
        norm_passes = [self.check_normality_full(g) for g in groups if len(g) >= 3]
        use_parametric = (sum(norm_passes) / len(norm_passes) > 0.95) if norm_passes else False
        
        if use_parametric:
            print(">> Testing system metrics via Standard Two-Way ANOVA Matrix.")
            formula = f"Q('{metric}') ~ C({factor_x}) * C({factor_hue})"
            model = ols(formula, data=clean_df).fit()
            anova_table = sm.stats.anova_lm(model, typ=2)
            p_inter = anova_table.loc[f"C({factor_x}):C({factor_hue})", "PR(>F)"]
        else:
            print(">> Testing system metrics via Non-Parametric Aligned Rank Transform (ART).")
            grand_mean = clean_df[metric].mean()
            mean_a = clean_df.groupby(factor_x)[metric].transform('mean')
            mean_b = clean_df.groupby(factor_hue)[metric].transform('mean')
            clean_df['art_aligned'] = clean_df[metric] - (mean_a - grand_mean) - (mean_b - grand_mean)
            clean_df['art_rank'] = stats.rankdata(clean_df['art_aligned'])
            
            formula = f"art_rank ~ C({factor_x}) * C({factor_hue})"
            model = ols(formula, data=clean_df).fit()
            anova_table = sm.stats.anova_lm(model, typ=2)
            p_inter = anova_table.loc[f"C({factor_x}):C({factor_hue})", "PR(>F)"]

        print(anova_table)
        print(f"\n>>> Interaction Significance Outcome: p = {p_inter:.4e}")
        
        if p_inter < 0.05:
            print("📢 RESULT: STATISTICALLY SIGNIFICANT INTERACTION (p < 0.05)")
        else:
            print("📢 RESULT: NOT STATISTICALLY SIGNIFICANT (p >= 0.05)")
        
        suff = f"_top{limit}" if limit else "_all"
        self._safe_append_sheet(anova_table.reset_index(), f"Int_{factor_x[:4]}_{factor_hue[:4]}{suff}")

        # --- STEP 2: APPLY THE VISUAL FILTER TO PRUNE THE PLOT ONLY ---
        plot_df = clean_df.copy()
        if limit is not None:
            is_error_metric = any(x in metric.upper() for x in ["MSE", "LOSS", "MAE"])
            
            if strategy == 'best':
                plot_agg = 'min' if is_error_metric else 'max'
            elif strategy == 'mean':
                plot_agg = 'mean'
            elif strategy == 'median':
                plot_agg = 'median'
            else: # strategy is None
                plot_agg = 'mean' if use_parametric else 'median'
            
            if plot_df[factor_x].nunique() > limit:
                top_x = plot_df.groupby(factor_x)[metric].agg(plot_agg).sort_values(ascending=is_error_metric).index[:limit].tolist()
                plot_df = plot_df[plot_df[factor_x].isin(top_x)].copy()
                plot_df[factor_x] = pd.Categorical(plot_df[factor_x], categories=top_x, ordered=True)
                print(f"🎯 [Visual Plot Filter] Pruned axis '{factor_x}' down to top {limit} choices based on {plot_agg.upper()}: {top_x}")
                
            if plot_df[factor_hue].nunique() > limit:
                top_hue = plot_df.groupby(factor_hue)[metric].agg(plot_agg).sort_values(ascending=is_error_metric).index[:limit].tolist()
                plot_df = plot_df[plot_df[factor_hue].isin(top_hue)].copy()
                print(f"🎯 [Visual Plot Filter] Pruned chart legends '{factor_hue}' down to top {limit} choices based on {plot_agg.upper()}: {top_hue}")

        plt.figure(figsize=(10, 5))
        sns.pointplot(x=factor_x, y=metric, hue=factor_hue, data=plot_df, capsize=.05, dodge=0.15)
        plt.title(f"Interaction Visual Map (Top Tier View Only)\nFull Model Significance: p = {p_inter:.4e}")
        plt.grid(True, linestyle=":", alpha=0.5)
        plt.xticks(rotation=15)
        plt.tight_layout()
        plt.show()

    def run_hyperparameter_correlation(self, factor_target, factor_context, metric='Val Global Open R2', strategy=None, filters=None):
        print(f"\n{'='*80}\nHYPERPARAMETER RANK CORRELATION: {factor_target} across {factor_context}")
        print(f"Metric Evaluated: {metric}\nFilters Applied: {filters}\n{'='*80}")
        
        self.df[metric] = pd.to_numeric(self.df[metric], errors='coerce')
        clean_df = self.df.dropna(subset=[metric]).copy()
        clean_df = self._apply_filters(clean_df, filters)
        
        if clean_df.empty:
            print("❌ Error: Filter criteria returned 0 tracking data entries.")
            return None
            
        is_error_metric = any(x in metric.upper() for x in ["MSE", "LOSS", "MAE"])
        
        if strategy == 'best':
            agg_metric = 'min' if is_error_metric else 'max'
        elif strategy == 'mean':
            agg_metric = 'mean'
        elif strategy == 'median':
            agg_metric = 'median'
        else: # strategy is None
            unique_targets = clean_df[factor_target].unique()
            norm_results = {idx: self.check_normality_full(clean_df[clean_df[factor_target]==idx][metric]) for idx in unique_targets}
            pct_normal = sum(norm_results.values()) / len(norm_results) if norm_results else 0
            use_parametric = pct_normal > 0.95
            agg_metric = 'mean' if use_parametric else 'median'
            
        pivot_df = clean_df.pivot_table(index=factor_target, columns=factor_context, values=metric, aggfunc=agg_metric)
        print(f"\n>>> Performance Profile Matrix ({agg_metric.upper()}):")
        print(pivot_df)
        
        if pivot_df.shape[1] > 1:
            corr_matrix = pivot_df.corr(method='spearman')
            print("\n>>> Spearman Rank Correlation Matrix between Contexts:")
            print(corr_matrix)
            self._safe_append_sheet(corr_matrix, f"CorrMat_{factor_target[:5]}_{factor_context[:5]}")
        else:
            print("\n⚠️ Warning: Only one environmental context variation found.")
            corr_matrix = None
            
        self._safe_append_sheet(pivot_df, f"CorrData_{factor_target[:5]}_{factor_context[:5]}")
        
        plt.figure(figsize=(10, 6))
        plot_data = pivot_df.T
        for option in plot_data.columns:
            plt.plot(plot_data.index, plot_data[option], marker='o', linewidth=2.5, label=str(option))
            
        plt.title(f"Hyperparameter Rank Profile consistency: {factor_target} vs {factor_context}\n({metric})")
        plt.xlabel(f"Environmental Context ({factor_context})")
        plt.ylabel(f"{agg_metric.upper()} {metric}")
        plt.grid(True, linestyle=":", alpha=0.6)
        plt.legend(title=factor_target, bbox_to_anchor=(1.05, 1), loc='upper left')
        plt.tight_layout()
        plt.show()
        
        return pivot_df, corr_matrix

    def run_cross_context_comparison(self, factor_target='map', context_col='initialization', context_1='serial', context_2='uniform', top_n=5, metric='Val Global Open R2', strategy=None, filters=None):
        print(f"\n{'='*80}\nCROSS-CONTEXT ELITE POOL COMPARISON: {factor_target} across {context_col}")
        print(f"Comparing context values: '{context_1}' vs '{context_2}'")
        print(f"Top N extracted from each: {top_n} | Base Filters applied: {filters}\n{'='*80}")
        
        self.df[metric] = pd.to_numeric(self.df[metric], errors='coerce')
        clean_df = self.df.dropna(subset=[metric]).copy()
        clean_df = self._apply_filters(clean_df, filters)
        
        if clean_df.empty:
            print("❌ Error: Filter parameters yielded an empty dataframe slice.")
            return None
            
        is_error_metric = any(x in metric.upper() for x in ["MSE", "LOSS", "MAE"])
        
        if strategy == 'best':
            plot_agg = 'min' if is_error_metric else 'max'
        elif strategy == 'mean':
            plot_agg = 'mean'
        elif strategy == 'median':
            plot_agg = 'median'
        else: # strategy is None
            unique_targets = clean_df[factor_target].unique()
            norm_results = {idx: self.check_normality_full(clean_df[clean_df[factor_target]==idx][metric]) for idx in unique_targets}
            pct_normal = sum(norm_results.values()) / len(norm_results) if norm_results else 0
            use_parametric = pct_normal > 0.95
            plot_agg = 'mean' if use_parametric else 'median'
        
        df_c1 = clean_df[clean_df[context_col].astype(str) == str(context_1)]
        if df_c1.empty: return None
        top_targets_1 = df_c1.groupby(factor_target)[metric].agg(plot_agg).sort_values(ascending=is_error_metric).index[:top_n].tolist()
        
        df_c2 = clean_df[clean_df[context_col].astype(str) == str(context_2)]
        if df_c2.empty: return None
        top_targets_2 = df_c2.groupby(factor_target)[metric].agg(plot_agg).sort_values(ascending=is_error_metric).index[:top_n].tolist()
        
        print(f"➔ Top {top_n} choices via {plot_agg.upper()} for context '{context_1}': {top_targets_1}")
        print(f"➔ Top {top_n} choices via {plot_agg.upper()} for context '{context_2}': {top_targets_2}")
        
        combined_targets = list(dict.fromkeys(top_targets_1 + top_targets_2))
        final_df = clean_df[(clean_df[factor_target].isin(combined_targets)) & (clean_df[context_col].isin([context_1, context_2]))].copy()
        final_df[factor_target] = pd.Categorical(final_df[factor_target], categories=combined_targets, ordered=True)
        
        summary_table = final_df.groupby([factor_target, context_col])[metric].agg(['mean', 'median', 'std', 'count', 'max', 'min']).unstack(level=1)
        self._safe_append_sheet(summary_table, f"Cross_{factor_target[:5]}_{context_1[:4]}_{context_2[:4]}")
        
        plt.figure(figsize=(12, 6))
        sns.boxplot(x=factor_target, y=metric, hue=context_col, data=final_df, palette=self.my_colors)
        sns.stripplot(x=factor_target, y=metric, hue=context_col, data=final_df, dodge=True, alpha=0.3, jitter=0.1, color='black')
        handles, labels = plt.gca().get_legend_handles_labels()
        plt.legend(handles[:2], labels[:2], title=context_col)
        plt.title(f"Elite Cross-Context Robustness Map ({metric})\nUnified Pool from Top {top_n} Performers ({plot_agg.upper()})")
        plt.grid(True, linestyle=":", alpha=0.5)
        plt.tight_layout()
        plt.show()
        
        return summary_table
    def run_cross_context_comparison(self, factor_target='map', context_col='initialization', context_1='serial', context_2='uniform', top_n=5, metric='Val Global Open R2', strategy=None, filters=None):
        """
        UPGRADE: Pools the top_n best hyperparameter variants from condition_1 and condition_2,
        removes overlapping duplicates, aggregates performance cross-wise, and outputs profiles.
        """
        print(f"\n{'='*80}\nCROSS-CONTEXT ELITE POOL COMPARISON: {factor_target} across {context_col}")
        print(f"Comparing context values: '{context_1}' vs '{context_2}'")
        print(f"Top N extracted from each: {top_n} | Base Filters applied: {filters}\n{'='*80}")
        
        self.df[metric] = pd.to_numeric(self.df[metric], errors='coerce')
        clean_df = self.df.dropna(subset=[metric]).copy()
        clean_df = self._apply_filters(clean_df, filters)
        
        if clean_df.empty:
            print("❌ Error: Filter parameters yielded an empty dataframe slice.")
            return None
            
        is_error_metric = any(x in metric.upper() for x in ["MSE", "LOSS", "MAE"])
        if strategy == 'best':
            plot_agg = 'min' if is_error_metric else 'max'
        elif strategy == 'mean':
            plot_agg = 'mean'
        elif strategy == 'median':
            plot_agg = 'median'
        else: # strategy is None
            unique_targets = clean_df[factor_target].unique()
            norm_results = {idx: self.check_normality_full(clean_df[clean_df[factor_target]==idx][metric]) for idx in unique_targets}
            pct_normal = sum(norm_results.values()) / len(norm_results) if norm_results else 0
            use_parametric = pct_normal > 0.95
            plot_agg = 'mean' if use_parametric else 'median'
        
        df_c1 = clean_df[clean_df[context_col].astype(str) == str(context_1)]
        if df_c1.empty: return None
        top_targets_1 = df_c1.groupby(factor_target)[metric].agg(plot_agg).sort_values(ascending=is_error_metric).index[:top_n].tolist()
        
        df_c2 = clean_df[clean_df[context_col].astype(str) == str(context_2)]
        if df_c2.empty: return None
        top_targets_2 = df_c2.groupby(factor_target)[metric].agg(plot_agg).sort_values(ascending=is_error_metric).index[:top_n].tolist()
        
        print(f"➔ Top {top_n} choices via {plot_agg.upper()} for context '{context_1}': {top_targets_1}")
        print(f"➔ Top {top_n} choices via {plot_agg.upper()} for context '{context_2}': {top_targets_2}")
        
        # 3. COLLAPSE OVERLAPS: Order preserving unique merge (Your exact deduplication logic)
        combined_targets = list(dict.fromkeys(top_targets_1 + top_targets_2))
        print(f"➔ Combined Elite Test Pool: {len(combined_targets)} unique configuration choices (Reduced from {top_n*2})")
        
        # 4. Isolate results for this unified elite list across BOTH conditions
        final_df = clean_df[
            (clean_df[factor_target].isin(combined_targets)) & 
            (clean_df[context_col].isin([context_1, context_2]))
        ].copy()
        
        # Sort categorical entries according to the discovered ranks
        final_df[factor_target] = pd.Categorical(final_df[factor_target], categories=combined_targets, ordered=True)
        
        # 5. Extract cross-wise aggregate results
        summary_table = final_df.groupby([factor_target, context_col])[metric].agg(['mean', 'median', 'std', 'count']).unstack(level=1)
        print("\n>>> Robust Unified Mean & Performance Cross-Matrix Table:")
        print(summary_table)
        
        self._safe_append_sheet(summary_table, f"Cross_{factor_target[:5]}_{context_1[:4]}_{context_2[:4]}")
        
        # 6. High-quality Comparative Plot Layout
        plt.figure(figsize=(12, 6))
        sns.boxplot(x=factor_target, y=metric, hue=context_col, data=final_df, palette=self.my_colors)
        sns.stripplot(x=factor_target, y=metric, hue=context_col, data=final_df, dodge=True, alpha=0.3, jitter=0.1, color='black')
        
        # Deduplicate categorical trace elements in the graph legend
        handles, labels = plt.gca().get_legend_handles_labels()
        plt.legend(handles[:2], labels[:2], title=context_col)
        
        plt.title(f"Elite Cross-Context Robustness Map ({metric})\nUnified Pool from Top {top_n} Performers")
        plt.grid(True, linestyle=":", alpha=0.5)
        plt.tight_layout()
        plt.show()
        
        return summary_table

# %% [markdown]
# # Interactive Notebook Execution Playground
# Run the cells below to execute individual analytical components.

# %%
# 1. INITIALIZE SYSTEM LAYER RUNTIME
# Replace 'your_data_file.xlsx' with the actual path to your dataset.
analyzer = InteractiveQNNAnalyzer("logs\\experiments_systematic\\qrc\\1_head_model_massive_search\\1_head_model.xlsx", save_enabled=True)

# %%
# 2. RUN EXPERIMENT CONFIGURATION STUDY (Global Perspective)
analyzer.run_study(group_by='configuration', metric='Val Global Open R2', filters={'initialization': 'identity'})
analyzer.run_study(group_by='configuration', metric='Val Global Open R2', filters={'initialization': 'uniform', 'encoding': 'serial'})


# %%
# 3. RUN HYPERPARAMETER FILTER SAMPLES (Filtering map variants within 1-head environments)
analyzer.run_study(group_by='map', metric='Val Global Open R2', filters={'encoding': "compact"}, strategy='best')

# %%
# 4. PAIRWISE EVALUATION MATRIX FOR TOP CONFIGURATIONS
analyzer.run_pairwise_analysis(group_by='configuration', metric='Val Global Open R2', limit=20, strategy='best', filters={'initialization': 'identity'})

# %%
# 5. TEST CO-DEPENDENCY INTERACTION SCHEMAS
analyzer.run_interaction_study(factor_x='encoding', factor_hue='map', metric='Val Global Open R2', limit = 5, strategy='best', filters={'initialization': 'uniform'})
# %%
# 5. TEST CO-DEPENDENCY INTERACTION SCHEMAS
analyzer.run_interaction_study(factor_x='initialization', factor_hue='map', metric='Val Global Open R2', limit = 5, strategy='best', filters={'encoding': 'serial'})
# %%
analyzer.run_hyperparameter_correlation(
    factor_target='map', 
    factor_context='initialization', 
    metric='Val Global Open R2', strategy='best',
    filters={'encoding': 'serial'}
)
# %%
my_filters = {'initialization': 'uniform'}

# Run the cross-evaluation engine
analyzer.run_cross_context_comparison(
    factor_target='map',         # The choice we want to optimize (e.g. maps)
    context_col='initialization', # The dimension defining our two environments
    context_1='compact',           # Condition A
    context_2='serial',          # Condition B
    top_n=5,                      # Extract top 5 from each
    metric='Val Global Open R2',
    strategy='best',
    filters=my_filters
)